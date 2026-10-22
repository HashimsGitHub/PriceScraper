from copy import copy
from datetime import datetime, timezone
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any
from uuid import uuid4

import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.utils import get_column_letter, range_boundaries


TEMPLATE_PATH = Path(__file__).resolve().parent / "PrestaShop_Product_Entry_Template.xlsx"
PRODUCTS_START_ROW = 4
PRODUCT_COLUMNS = [
    "Selected",
    "Product Name *",
    "Product Type *",
    "Reference (SKU) *",
    "Brand",
    "Description (HTML) *",
    "Main Category *",
    "Quantity *",
    "Weight (lb) *",
    "Google Price (USD)",
    "Retail Price (tax excl.) *",
    "Tax Rule / State",
    "Publish Status *",
    "Merchant",
    "Google Shopping Link",
]
NUMERIC_COLUMNS = {
    "Quantity *",
    "Weight (lb) *",
    "Google Price (USD)",
    "Retail Price (tax excl.) *",
}
TEMPLATE_FIELD_MAP = {
    "Product Name *": "Product Name *",
    "Product Type *": "Product Type *",
    "Reference (SKU) *": "Reference (SKU) *",
    "Brand": "Brand",
    "Description (HTML) *": "Description (HTML) *",
    "Main Category *": "Main Category *",
    "Quantity *": "Quantity *",
    "Weight (lb) *": "Weight (lb) *",
    "Retail Price (tax excl.) *": "Retail Price (tax excl.) *",
    "Tax Rule *": "Tax Rule / State",
    "Publish Status *": "Publish Status *",
}

st.set_page_config(page_title="Product Price Scraper", page_icon="🛍️", layout="wide")


def get_template_options() -> dict[str, list[str]]:
    """Read import dropdown options from the supplied workbook."""
    workbook = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True, data_only=True)
    if "Products" not in workbook.sheetnames or "Lists" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("The template must contain Products and Lists worksheets.")

    lists_sheet = workbook["Lists"]

    def values_in_column(column: int) -> list[str]:
        return [
            str(row[0]).strip()
            for row in lists_sheet.iter_rows(
                min_row=2, min_col=column, max_col=column, values_only=True
            )
            if row[0] is not None and str(row[0]).strip()
        ]

    options = {
        "product_types": values_in_column(1),
        "tax_rules": values_in_column(5),
        "publish_statuses": values_in_column(8),
        "brands": values_in_column(10),
        "categories": values_in_column(11),
    }
    workbook.close()
    return options


def parse_price(value: Any) -> float | None:
    """Parse a Shopping price only when it is a recognizable currency amount."""
    if value is None:
        return None
    text = str(value).strip()
    match = re.search(r"(?:\$|USD\s*)\s*([\d,]+(?:\.\d{1,2})?)", text, re.IGNORECASE)
    if match is None:
        match = re.fullmatch(r"[\d,]+(?:\.\d{1,2})?", text)
        if match is None:
            return None
        amount = match.group(0)
    else:
        amount = match.group(1)
    return float(amount.replace(",", ""))


def fetch_google_shopping_prices(
    query: str, target_count: int, api_key: str
) -> list[dict[str, Any]]:
    """Fetch the complete Google Shopping result list returned by SerpApi."""
    response = requests.get(
        "https://serpapi.com/search.json",
        params={
            "engine": "google_shopping",
            "q": query,
            "location": "United States",
            "hl": "en",
            "gl": "us",
            "num": target_count,
            "api_key": api_key,
        },
        timeout=30,
    )
    response.raise_for_status()
    data = response.json()
    if not isinstance(data, dict):
        raise ValueError("SerpApi returned an invalid response.")
    if data.get("error"):
        raise ValueError(f"SerpApi error: {data['error']}")

    shopping_results = data.get("shopping_results", [])
    if not isinstance(shopping_results, list):
        raise ValueError("SerpApi returned an invalid shopping_results value.")
    if any(not isinstance(item, dict) for item in shopping_results):
        raise ValueError("SerpApi returned a shopping result that is not an object.")
    return shopping_results


def make_product_entry(item: dict[str, Any], tax_rule: str) -> dict[str, Any]:
    product_id = item.get("product_id")
    return {
        "Selected": False,
        "Product Name *": str(item.get("title") or ""),
        "Product Type *": "Standard product",
        "Reference (SKU) *": str(product_id) if product_id else "",
        "Brand": "",
        "Description (HTML) *": "",
        "Main Category *": "",
        "Quantity *": None,
        "Weight (lb) *": None,
        "Google Price (USD)": parse_price(item.get("price")),
        "Retail Price (tax excl.) *": parse_price(item.get("price")),
        "Tax Rule / State": tax_rule,
        "Publish Status *": "Offline",
        "Merchant": str(item.get("source") or ""),
        "Google Shopping Link": str(item.get("product_link") or item.get("link") or ""),
    }


def normalize_editor_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized = []
    for record in records:
        clean_record = {}
        for column in PRODUCT_COLUMNS:
            value = record.get(column)
            if pd.isna(value):
                value = None if column in NUMERIC_COLUMNS else ""
            clean_record[column] = value
        normalized.append(clean_record)
    return normalized


def validate_selected_products(
    products: list[dict[str, Any]],
    options: dict[str, list[str]],
    existing_references: set[str] | None = None,
) -> list[str]:
    errors = []
    used_references = set(existing_references or ())

    for row_number, product in enumerate(products, start=1):
        prefix = f"Selected item {row_number}"
        name = str(product.get("Product Name *") or "").strip()
        if not name:
            errors.append(f"{prefix}: Product Name is required.")
        elif len(name) > 128 or re.search(r"[\^<>=\{\}]", name):
            errors.append(
                f"{prefix}: Product Name must be at most 128 characters and avoid ^, <, >, =, {{, and }}."
            )

        product_type = str(product.get("Product Type *") or "").strip()
        if product_type not in options["product_types"]:
            errors.append(f"{prefix}: choose a Product Type listed on the template's Lists sheet.")

        reference = str(product.get("Reference (SKU) *") or "").strip()
        if not reference:
            errors.append(f"{prefix}: Reference (SKU) is required.")
        elif len(reference) > 64:
            errors.append(f"{prefix}: Reference (SKU) must be at most 64 characters.")
        elif reference in used_references:
            errors.append(f"{prefix}: Reference (SKU) {reference!r} already exists in this export.")
        else:
            used_references.add(reference)

        if not str(product.get("Description (HTML) *") or "").strip():
            errors.append(f"{prefix}: Description (HTML) is required.")

        category = str(product.get("Main Category *") or "").strip()
        if category not in options["categories"]:
            errors.append(f"{prefix}: Main Category must match a value on the template's Lists sheet.")

        quantity = product.get("Quantity *")
        try:
            quantity_value = float(quantity)
            if not quantity_value.is_integer() or quantity_value < 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"{prefix}: Quantity must be a whole number greater than or equal to 0.")

        for field in ("Weight (lb) *", "Retail Price (tax excl.) *"):
            try:
                numeric_value = float(product.get(field))
                if not math.isfinite(numeric_value) or numeric_value < 0:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append(f"{prefix}: {field} must be a number greater than or equal to 0.")

        tax_rule = str(product.get("Tax Rule / State") or "").strip()
        if tax_rule not in options["tax_rules"]:
            errors.append(f"{prefix}: Tax Rule / State must match a value on the template's Lists sheet.")

        brand = str(product.get("Brand") or "").strip()
        if brand and brand not in options["brands"]:
            errors.append(f"{prefix}: Brand must exactly match a value on the template's Lists sheet.")

        publish_status = str(product.get("Publish Status *") or "").strip()
        if publish_status not in options["publish_statuses"]:
            errors.append(f"{prefix}: Publish Status must match a value on the template's Lists sheet.")

    return errors


def existing_product_references() -> set[str]:
    workbook = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True, data_only=True)
    products_sheet = workbook["Products"]
    headers = [cell.value for cell in products_sheet[2]]
    try:
        reference_column = headers.index("Reference (SKU) *") + 1
    except ValueError as error:
        workbook.close()
        raise ValueError("The Products sheet is missing the Reference (SKU) * column.") from error

    references = {
        str(row[0]).strip()
        for row in products_sheet.iter_rows(
            min_row=PRODUCTS_START_ROW,
            min_col=reference_column,
            max_col=reference_column,
            values_only=True,
        )
        if row[0] is not None and str(row[0]).strip()
    }
    workbook.close()
    return references


def build_prestashop_workbook(products: list[dict[str, Any]]) -> bytes:
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    worksheet = workbook["Products"]
    headers = [cell.value for cell in worksheet[2]]
    if not headers or any(not header for header in headers):
        raise ValueError("The Products sheet has a missing header in row 2.")

    header_columns = {str(header): index + 1 for index, header in enumerate(headers)}
    missing_fields = set(TEMPLATE_FIELD_MAP) - set(header_columns)
    if missing_fields:
        raise ValueError(
            "The Products sheet is missing required columns: "
            + ", ".join(sorted(missing_fields))
        )

    last_data_row = PRODUCTS_START_ROW - 1
    for row_number in range(PRODUCTS_START_ROW, worksheet.max_row + 1):
        if any(
            worksheet.cell(row_number, column).value not in (None, "")
            for column in range(1, len(headers) + 1)
        ):
            last_data_row = row_number
    first_output_row = last_data_row + 1
    original_max_row = worksheet.max_row

    for offset, product in enumerate(products):
        row_number = first_output_row + offset
        if row_number > original_max_row:
            for column in range(1, len(headers) + 1):
                source = worksheet.cell(PRODUCTS_START_ROW, column)
                destination = worksheet.cell(row_number, column)
                if source.has_style:
                    destination._style = copy(source._style)
                if source.number_format:
                    destination.number_format = source.number_format
            source_height = worksheet.row_dimensions[PRODUCTS_START_ROW].height
            if source_height is not None:
                worksheet.row_dimensions[row_number].height = source_height

        values = {
            template_field: product.get(entry_field)
            for template_field, entry_field in TEMPLATE_FIELD_MAP.items()
        }
        values["MPN"] = product.get("Reference (SKU) *")
        for header, column in header_columns.items():
            cell = worksheet.cell(row_number, column, values.get(header))
            if isinstance(cell.value, str):
                cell.data_type = "s"

    last_output_row = first_output_row + len(products) - 1
    for validation in worksheet.data_validations.dataValidation:
        for cell_range in list(validation.ranges):
            min_col, min_row, max_col, max_row = range_boundaries(str(cell_range))
            if min_row <= last_output_row and max_row < last_output_row:
                validation.add(
                    f"{get_column_letter(min_col)}{max_row + 1}:"
                    f"{get_column_letter(max_col)}{last_output_row}"
                )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def build_search_results_workbook(search_runs: list[dict[str, Any]]) -> bytes:
    workbook = openpyxl.Workbook()
    results_sheet = workbook.active
    results_sheet.title = "All Search Results"
    runs_sheet = workbook.create_sheet("Search Runs")

    metadata_columns = ["Search ID", "Search Query", "Searched At (UTC)", "Initial Tax Rule / State"]
    raw_columns = []
    for run in search_runs:
        for result in run["raw_results"]:
            for column in result:
                if column not in raw_columns and column not in metadata_columns:
                    raw_columns.append(column)

    results_sheet.append(metadata_columns + raw_columns)
    for cell in results_sheet[1]:
        cell.data_type = "s"
    for run in search_runs:
        for result in run["raw_results"]:
            row = [
                run["id"],
                run["query"],
                run["searched_at"],
                run["tax_rule"],
            ]
            row.extend(excel_value(result.get(column)) for column in raw_columns)
            results_sheet.append(row)
            for cell in results_sheet[results_sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    runs_sheet.append(["Search ID", "Search Query", "Searched At (UTC)", "Tax Rule / State", "Result Count"])
    for cell in runs_sheet[1]:
        cell.data_type = "s"
    for run in search_runs:
        runs_sheet.append(
            [
                run["id"],
                run["query"],
                run["searched_at"],
                run["tax_rule"],
                len(run["raw_results"]),
            ]
        )
        for cell in runs_sheet[runs_sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for sheet in (results_sheet, runs_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            header_font = copy(cell.font)
            header_font.bold = True
            cell.font = header_font
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 14), 50)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


if not TEMPLATE_PATH.is_file():
    st.error(
        f"Template workbook not found: {TEMPLATE_PATH.name}. "
        "Place it beside app.py before starting the app."
    )
    st.stop()

template_options = get_template_options()
if not template_options["tax_rules"]:
    st.error("No Tax Rule options were found in the template's Lists sheet.")
    st.stop()

st.title("Product Price Scraper")
st.markdown(
    "Search Google Shopping, review listings, and select products to append to the "
    "PrestaShop product-entry template."
)

api_key = st.secrets.get("SERPAPI_KEY")
with st.sidebar:
    st.header("Settings")
    if api_key:
        st.success("SerpApi key loaded from Streamlit Secrets.")
    else:
        api_key = st.text_input("SerpApi key", type="password")
    st.caption("Searches use Google Shopping results for the United States.")

if "search_runs" not in st.session_state:
    st.session_state.search_runs = []

with st.form("product_search"):
    search_query = st.text_input("Product search", placeholder="e.g. wireless barcode scanner")
    max_results = st.slider("Maximum results per search", min_value=5, max_value=50, value=15, step=5)
    default_tax_index = (
        template_options["tax_rules"].index("By State")
        if "By State" in template_options["tax_rules"]
        else 0
    )
    tax_rule = st.selectbox(
        "Tax Rule / State",
        options=template_options["tax_rules"],
        index=default_tax_index,
        help="This value is copied to each result and can be changed per product before export.",
    )
    search_submitted = st.form_submit_button("Search Google Shopping", type="primary")

if search_submitted:
    if not search_query.strip():
        st.error("Enter a product to search for.")
    elif not api_key:
        st.error("Configure SERPAPI_KEY in Streamlit Secrets or enter a key in the sidebar.")
    else:
        with st.spinner("Searching Google Shopping via SerpApi..."):
            try:
                raw_results = fetch_google_shopping_prices(
                    search_query.strip(), max_results, api_key
                )
            except (requests.RequestException, ValueError) as error:
                st.error(f"Search failed: {error}")
            else:
                run = {
                    "id": str(uuid4()),
                    "query": search_query.strip(),
                    "searched_at": datetime.now(timezone.utc).isoformat(),
                    "tax_rule": tax_rule,
                    "raw_results": raw_results,
                    "entries": [
                        make_product_entry(item, tax_rule) for item in raw_results
                    ],
                }
                st.session_state.search_runs.append(run)
                st.success(f"Added {len(raw_results)} result(s) for “{run['query']}”.")

search_runs = st.session_state.search_runs
if search_runs:
    total_results = sum(len(run["raw_results"]) for run in search_runs)
    st.subheader(f"Search Results ({total_results} across {len(search_runs)} search(es))")
    st.info(
        "Review each selected item before export. Quantity, weight, description, and category "
        "must be supplied; the Shopping price may include tax, while the template expects "
        "a tax-exclusive retail price."
    )

    for run_number, run in enumerate(search_runs, start=1):
        with st.expander(f"Search {run_number}: {run['query']} ({len(run['raw_results'])} results)", expanded=True):
            st.caption(
                f"Searched {run['searched_at']} · Initial Tax Rule / State: {run['tax_rule']}"
            )
            edited = st.data_editor(
                pd.DataFrame(run["entries"], columns=PRODUCT_COLUMNS),
                key=f"product_editor_{run['id']}",
                hide_index=True,
                use_container_width=True,
                num_rows="fixed",
                column_order=PRODUCT_COLUMNS,
                column_config={
                    "Selected": st.column_config.CheckboxColumn("Import", default=False),
                    "Product Name *": st.column_config.TextColumn(
                        "Product Name *", help="Pre-filled from the Shopping listing."
                    ),
                    "Product Type *": st.column_config.TextColumn(
                        "Product Type *", help="Must match the template's Lists sheet."
                    ),
                    "Reference (SKU) *": st.column_config.TextColumn(
                        "Reference (SKU) *",
                        help="Google product_id is only a candidate; verify the manufacturer's SKU.",
                    ),
                    "Brand": st.column_config.TextColumn(
                        "Brand", help="If entered, it must exactly match a brand in the template."
                    ),
                    "Description (HTML) *": st.column_config.TextColumn(
                        "Description (HTML) *", help="Required by the PrestaShop template."
                    ),
                    "Main Category *": st.column_config.TextColumn(
                        "Main Category *",
                        help="Must exactly match a category in the template's Lists sheet.",
                    ),
                    "Quantity *": st.column_config.NumberColumn(
                        "Quantity *", min_value=0, step=1, format="%d"
                    ),
                    "Weight (lb) *": st.column_config.NumberColumn(
                        "Weight (lb) *", min_value=0, step=0.01
                    ),
                    "Google Price (USD)": st.column_config.NumberColumn(
                        "Google Price (USD)", disabled=True, format="$%.2f"
                    ),
                    "Retail Price (tax excl.) *": st.column_config.NumberColumn(
                        "Retail Price (tax excl.) *", min_value=0, step=0.01
                    ),
                    "Tax Rule / State": st.column_config.TextColumn(
                        "Tax Rule / State",
                        help="Must match a tax rule in the template; change this per item if needed.",
                    ),
                    "Publish Status *": st.column_config.TextColumn(
                        "Publish Status *", help="Defaults to Offline for review before publishing."
                    ),
                    "Merchant": st.column_config.TextColumn("Merchant", disabled=True),
                    "Google Shopping Link": st.column_config.TextColumn(
                        "Google Shopping Link", disabled=True
                    ),
                },
            )
            run["entries"] = normalize_editor_records(edited.to_dict("records"))

    selected_products = [
        entry
        for run in search_runs
        for entry in run["entries"]
        if bool(entry.get("Selected"))
    ]
    st.write(f"Selected for PrestaShop: **{len(selected_products)}**")

    results_excel = build_search_results_workbook(search_runs)
    st.download_button(
        "Download all search results (.xlsx)",
        data=results_excel,
        file_name="PrestaShop_Search_Results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    if selected_products:
        validation_errors = validate_selected_products(
            selected_products, template_options, existing_product_references()
        )
        if validation_errors:
            st.error("Complete or correct the selected items before downloading:")
            for error in validation_errors:
                st.write(f"- {error}")
        else:
            prestashop_excel = build_prestashop_workbook(selected_products)
            st.download_button(
                "Download PrestaShop_Product_Entry_Data.xlsx",
                data=prestashop_excel,
                file_name="PrestaShop_Product_Entry_Data.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )
    else:
        st.caption("Select one or more results and complete the required product fields to export.")
else:
    st.caption("Run a search to view product listings here. You can perform multiple searches in one session.")
